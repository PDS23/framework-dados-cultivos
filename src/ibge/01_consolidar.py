# -*- coding: utf-8 -*-
"""
Estágio 1/6 do pipeline IBGE: consolida as planilhas brutas exportadas do
SIDRA (Tabela 1613 - Produção Agrícola Municipal) numa única planilha por
cultura, com 1 aba por variável e todos os anos disponíveis em colunas.

Entrada esperada: data/brutos/ibge/<cultura>/*.xlsx — um ou mais arquivos
exportados do SIDRA PARA ESSA CULTURA (se o período completo não couber
numa exportação só, baixe em pedaços e coloque todos os arquivos aqui; o
script funde por ano). Cada arquivo deve ter as 9 abas padrão do SIDRA:
  Área destinada à colheita (Hectares)
  Área destinada à colheita - percentual do total geral
  Área colhida (Hectares)
  Área colhida - percentual do total geral
  Quantidade produzida (Toneladas)
  Rendimento médio da produção
  Valor da produção
  Valor da produção - percentual do total geral
  Notas
E, em cada aba, linhas por Nível (Brasil/UF/Município) + Código + Nome,
com 1 coluna por ano.

Saída: data/ibge/consolidado/<cultura>.xlsx (mesmas 9 abas, 1 arquivo só,
todos os anos baixados mesclados).

Uso: python src/ibge/01_consolidar.py <cultura>
"""
import sys
import glob
import os
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "comum"))
from config import paths_ibge, cultura as get_cultura  # noqa: E402

VAR_SHEET_NAMES = [
    "Área destinada à colheita (Hectares)",
    "Área destinada à colheita - percentual do total geral",
    "Área colhida (Hectares)",
    "Área colhida - percentual do total geral",
    "Quantidade produzida (Toneladas)",
    "Rendimento médio da produção",
    "Valor da produção",
    "Valor da produção - percentual do total geral",
]


def consolida(cultura_slug):
    get_cultura(cultura_slug)  # valida que a cultura existe no cadastro
    paths = paths_ibge(cultura_slug)
    brutos_dir = paths["brutos"]
    out_path = paths["consolidado"]
    out_path.parent.mkdir(parents=True, exist_ok=True)

    arquivos = sorted(glob.glob(str(brutos_dir / "*.xlsx")))
    if not arquivos:
        raise FileNotFoundError(
            f"Nenhum .xlsx encontrado em {brutos_dir}. Baixe a exportação do SIDRA "
            f"(Tabela 1613, produto = cultura, níveis Brasil/UF/Município) e coloque aqui."
        )
    print(f"[{cultura_slug}] Arquivos brutos encontrados: {len(arquivos)}")
    for f in arquivos:
        print("  -", f)

    # data[var_idx][codigo] = {"nome":..., "anos": {ano: valor}}
    data = [dict() for _ in VAR_SHEET_NAMES]
    all_years = set()
    notas_text = None

    for fpath in arquivos:
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        # Identifica a aba de cada variável pela POSIÇÃO (mesma ordem sempre,
        # 1 aba por variável + "Notas" no final), não pelo nome — o Excel
        # trunca nomes de aba em 31 caracteres, e os nomes completos das
        # variáveis do SIDRA passam disso, então o nome salvo no arquivo já
        # pode vir cortado (com ou sem "...").
        for var_idx, _nome_var in enumerate(VAR_SHEET_NAMES):
            if var_idx >= len(wb.sheetnames):
                continue
            ws = wb[wb.sheetnames[var_idx]]
            todas_linhas = list(ws.iter_rows(values_only=True))

            # Formato do cabeçalho varia (exportação direta do SIDRA tem ~5
            # linhas de cabeçalho com título/variável/rótulos; um export já
            # limpo pode ter só 1 linha "Nivel,Codigo,Nome,ano1,ano2,..."). Em
            # vez de assumir um número fixo de linhas, detecta automaticamente:
            # a primeira linha de dados é a primeira cujo col[0] é "UF"/"MU"/"BR";
            # a linha de anos é a linha ANTERIOR a essa com mais células que
            # parecem um ano de 4 dígitos.
            idx_primeira_linha_dado = next(
                (i for i, row in enumerate(todas_linhas) if row and row[0] in ("UF", "MU", "BR")),
                None,
            )
            if idx_primeira_linha_dado is None:
                continue  # aba sem dados reconhecíveis

            def conta_anos(row):
                if row is None:
                    return 0
                n = 0
                for val in row:
                    if val not in (None, ""):
                        digits = "".join(ch for ch in str(val)[:4] if ch.isdigit())
                        if len(digits) == 4:
                            n += 1
                return n

            candidatas = todas_linhas[:idx_primeira_linha_dado]
            year_row = max(candidatas, key=conta_anos) if candidatas else []

            col_year = {}
            for ci, val in enumerate(year_row):
                if ci < 3:
                    continue
                if val not in (None, ""):
                    digits = "".join(ch for ch in str(val)[:4] if ch.isdigit())
                    if len(digits) == 4:
                        col_year[ci] = digits

            for row in todas_linhas[idx_primeira_linha_dado:]:
                if row is None or len(row) < 3:
                    continue
                nivel, codigo, nome = row[0], row[1], row[2]
                if nivel not in ("UF", "MU"):
                    continue
                codigo_key = str(codigo) if codigo is not None else None
                if codigo_key is None:
                    continue
                for ci, ano in col_year.items():
                    if ci >= len(row):
                        continue
                    val = row[ci]
                    all_years.add(ano)
                    entry = data[var_idx].setdefault(codigo_key, {"nome": nome, "nivel": nivel, "anos": {}})
                    entry["anos"][ano] = val
        if "Notas" in wb.sheetnames:
            notas_text = list(wb["Notas"].iter_rows(values_only=True))
        wb.close()

    if not all_years:
        raise ValueError(f"Nenhum ano lido dos arquivos de {cultura_slug} — confira o formato das planilhas.")
    years_sorted = sorted(all_years, key=lambda y: int(y))
    print(f"[{cultura_slug}] Anos consolidados: {years_sorted[0]}-{years_sorted[-1]} ({len(years_sorted)} anos)")

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)
    for var_idx, var_name in enumerate(VAR_SHEET_NAMES):
        ws = out_wb.create_sheet(title=var_name[:31])
        ws.append(["Nivel", "Codigo", "Nome"] + years_sorted)
        for codigo, entry in sorted(data[var_idx].items(), key=lambda kv: (kv[1]["nivel"], kv[0])):
            row = [entry["nivel"], codigo, entry["nome"]] + [entry["anos"].get(a) for a in years_sorted]
            ws.append(row)
        ws.freeze_panes = "D2"
        ws.column_dimensions["C"].width = 30

    if notas_text:
        ws_notas = out_wb.create_sheet(title="Notas")
        for row in notas_text:
            ws_notas.append(list(row) if row else [])
        ws_notas.column_dimensions["A"].width = 100

    out_wb.save(out_path)
    print(f"[{cultura_slug}] Consolidado salvo em: {out_path}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python 01_consolidar.py <cultura>  (slug de config/culturas.yaml)")
        sys.exit(1)
    consolida(sys.argv[1])
