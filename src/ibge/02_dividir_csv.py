# -*- coding: utf-8 -*-
"""
Estágio 2/6 do pipeline IBGE: lê o consolidado (data/ibge/consolidado/<cultura>.xlsx)
e gera, para cada variável, um CSV separado em nível de MUNICÍPIO:
  data/ibge/processado/<cultura>/<variavel>.csv

Colunas de saída: Codigo, Nome, Teve_Prod, <ano1>, <ano2>, ...
"-" (zero absoluto do SIDRA) e "..." (não informado) viram 0 numérico.
Teve_Prod = 1 se algum ano tiver valor > 0.

Uso: python src/ibge/02_dividir_csv.py <cultura>
"""
import sys
import os
import csv
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "comum"))
from config import paths_ibge  # noqa: E402

VALORES_AUSENTES = (None, "", "-", "...")

# Mesma ORDEM de VAR_SHEET_NAMES em 01_consolidar.py. Identificação por
# posição, não por nome da aba — nomes de variável do SIDRA passam de 31
# caracteres e o Excel trunca o título da aba nesse limite, o que quebraria
# um match por nome exato entre os dois estágios.
VAR_SLUGS = [
    "area_destinada_colheita_ha",
    "area_destinada_colheita_pct",
    "area_colhida_ha",
    "area_colhida_pct",
    "quantidade_produzida_t",
    "rendimento_medio_kg_ha",
    "valor_producao",
    "valor_producao_pct",
]


def trata_valor(v):
    return 0 if v in VALORES_AUSENTES else v


def tem_dado_positivo(valores):
    for v in valores:
        try:
            if float(v) > 0:
                return 1
        except (TypeError, ValueError):
            continue
    return 0


def divide(cultura_slug):
    paths = paths_ibge(cultura_slug)
    src_path = paths["consolidado"]
    out_dir = paths["processado"]
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(src_path, read_only=True, data_only=True)
    count = 0
    for idx_var, slug in enumerate(VAR_SLUGS):
        if idx_var >= len(wb.sheetnames):
            continue
        ws = wb[wb.sheetnames[idx_var]]
        out_path = out_dir / f"{slug}.csv"
        with open(out_path, "w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.writer(fh)
            header_escrito = False
            anos_idx = None
            for row in ws.iter_rows(values_only=True):
                if row is None:
                    continue
                nivel, codigo, nome = row[0], row[1], row[2]
                if not header_escrito:
                    anos = row[3:]
                    writer.writerow(["Codigo", "Nome", "Teve_Prod"] + list(anos))
                    header_escrito = True
                    continue
                if nivel != "MU":
                    continue
                anos = [trata_valor(v) for v in row[3:]]
                flag = tem_dado_positivo(anos)
                writer.writerow([codigo, nome, flag] + anos)
        count += 1
        print(f"[{cultura_slug}] {slug}.csv salvo em {out_path}")
    wb.close()
    print(f"[{cultura_slug}] Total de CSVs gerados: {count}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python 02_dividir_csv.py <cultura>")
        sys.exit(1)
    divide(sys.argv[1])
