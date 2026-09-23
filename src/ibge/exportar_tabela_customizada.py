# -*- coding: utf-8 -*-
"""
Exportador "leve": pega um arquivo bruto do SIDRA (qualquer recorte — Brasil
inteiro ou só algumas UFs) e gera uma tabela larga (1 linha por município,
1 coluna por variável x ano) em .xlsx e .gpkg, sem passar pelo pipeline
completo (histórico nacional + recorte regional + mapas). Útil para recortes
pontuais: poucas variáveis, poucos anos, ou uma cultura só disponível para
um estado.

Colunas de saída: Codigo, Nome, UF, <slug_variavel>_<ano>, ...

Geometria do .gpkg vem da malha municipal em data/geo/ (mesma usada pelo
resto do framework — rode src/ibge/03_gerar_geojson.py para qualquer
cultura pelo menos 1x antes, se ainda não tiver a malha).

Uso:
    python src/ibge/exportar_tabela_customizada.py \
        --bruto "caminho/para/tabela.xlsx" \
        --variaveis "0:area_destinada_colheita_ha,4:quantidade_produzida_t" \
        --anos 2023,2024,2025 \
        --uf RS \
        --saida "caminho/para/mate_rs_2023_2025"

--variaveis: "<posição da aba>:<slug>", separados por vírgula. A posição é
o índice (0-based) da aba na planilha bruta, na ordem padrão do SIDRA:
  0 = Área destinada à colheita (ha)      4 = Quantidade produzida (t)
  1 = Área destinada à colheita (%)       5 = Rendimento médio (kg/ha)
  2 = Área colhida (ha)                   6 = Valor da produção
  3 = Área colhida (%)                    7 = Valor da produção (%)

--saida: caminho sem extensão — gera "<saida>.xlsx" e "<saida>.gpkg". Se
omitido, salva ao lado do arquivo bruto.
"""
import argparse
import os
import sys
import openpyxl
import pandas as pd
import geopandas as gpd

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "comum"))
from config import DATA  # noqa: E402
from sidra_parsing import detecta_cabecalho, separa_nome_uf, trata_valor  # noqa: E402

MALHA_SIMPLIFICADA = DATA / "geo" / "municipios_brasil_simplificado.geojson"
MALHA_COMPLETA = DATA / "geo" / "malha_municipios_brasil.shp"


def parse_variaveis(spec):
    pares = []
    for item in spec.split(","):
        idx_str, slug = item.split(":")
        pares.append((int(idx_str.strip()), slug.strip()))
    return pares


def exporta(bruto_path, variaveis, anos, uf_filtro, saida_base):
    anos = [str(a) for a in anos]
    wb = openpyxl.load_workbook(bruto_path, read_only=True, data_only=True)

    # {(codigo): {"nome":..., "uf":..., valores: {(slug,ano): valor}}}
    municipios = {}

    for idx_aba, slug in variaveis:
        if idx_aba >= len(wb.sheetnames):
            raise ValueError(f"Aba de índice {idx_aba} não existe neste arquivo (só tem {len(wb.sheetnames)} abas).")
        ws = wb[wb.sheetnames[idx_aba]]
        todas_linhas, idx_primeira_linha_dado, col_ano = detecta_cabecalho(ws)
        if idx_primeira_linha_dado is None:
            raise ValueError(f"Não encontrei linhas de dados (MU/UF/BR) na aba {idx_aba} ({ws.title}).")

        cols_pedidas = {ci: ano for ci, ano in col_ano.items() if ano in anos}
        if not cols_pedidas:
            anos_disponiveis = sorted(set(col_ano.values()))
            raise ValueError(f"Nenhum dos anos {anos} está na aba {idx_aba} ({ws.title}). Anos disponíveis: {anos_disponiveis}")

        for row in todas_linhas[idx_primeira_linha_dado:]:
            if row is None or len(row) < 3 or row[0] != "MU":
                continue
            codigo, municipio_raw = str(row[1]), row[2]
            nome, uf = separa_nome_uf(municipio_raw)
            if uf_filtro and uf != uf_filtro:
                continue
            entry = municipios.setdefault(codigo, {"nome": nome, "uf": uf, "valores": {}})
            for ci, ano in cols_pedidas.items():
                if ci < len(row):
                    entry["valores"][(slug, ano)] = trata_valor(row[ci])
    wb.close()

    if not municipios:
        raise ValueError(f"Nenhum município encontrado (UF filtrada: {uf_filtro!r}) — confira --uf e o conteúdo do arquivo bruto.")

    colunas_valor = [f"{slug}_{ano}" for _idx, slug in variaveis for ano in anos]
    linhas = []
    for codigo, entry in sorted(municipios.items()):
        linha = {"Codigo": codigo, "Nome": entry["nome"], "UF": entry["uf"]}
        for _idx, slug in variaveis:
            for ano in anos:
                linha[f"{slug}_{ano}"] = entry["valores"].get((slug, ano))
        linhas.append(linha)
    df = pd.DataFrame(linhas, columns=["Codigo", "Nome", "UF"] + colunas_valor)
    for col in colunas_valor:
        df[col] = df[col].astype("Int64")  # inteiro com suporte a nulo (valores "..." viram <NA>)

    saida_base = str(saida_base)
    xlsx_path = saida_base + ".xlsx"
    gpkg_path = saida_base + ".gpkg"

    df.to_excel(xlsx_path, index=False)
    print(f"Excel salvo em: {xlsx_path} ({len(df)} municípios)")

    malha_path = MALHA_SIMPLIFICADA if MALHA_SIMPLIFICADA.exists() else MALHA_COMPLETA
    if not malha_path.exists():
        print(f"Aviso: malha municipal não encontrada ({MALHA_SIMPLIFICADA} nem {MALHA_COMPLETA}) — pulando o .gpkg.")
        return
    malha = gpd.read_file(malha_path)
    col_codigo_malha = "CD_MUN" if "CD_MUN" in malha.columns else "Codigo"
    malha[col_codigo_malha] = malha[col_codigo_malha].astype(str)
    malha = malha.rename(columns={col_codigo_malha: "Codigo"})[["Codigo", "geometry"]]

    gdf = malha.merge(df, on="Codigo", how="inner")
    faltantes = set(df["Codigo"]) - set(gdf["Codigo"])
    if faltantes:
        print(f"Aviso: {len(faltantes)} município(s) sem geometria na malha (fora do .gpkg): {sorted(faltantes)[:5]}...")
    gdf.to_file(gpkg_path, driver="GPKG")
    print(f"GeoPackage salvo em: {gpkg_path} ({len(gdf)} municípios)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--bruto", required=True, help="caminho do .xlsx bruto do SIDRA")
    parser.add_argument("--variaveis", required=True, help='"<posição da aba>:<slug>,..." ex.: "0:area_destinada_colheita_ha,4:quantidade_produzida_t"')
    parser.add_argument("--anos", required=True, help="anos separados por vírgula, ex.: 2023,2024,2025")
    parser.add_argument("--uf", default=None, help="sigla da UF pra filtrar (opcional)")
    parser.add_argument("--saida", default=None, help="caminho de saída sem extensão (default: ao lado do bruto)")
    args = parser.parse_args()

    variaveis = parse_variaveis(args.variaveis)
    anos = [int(a) for a in args.anos.split(",")]
    saida = args.saida or os.path.splitext(args.bruto)[0] + "_exportado"

    exporta(args.bruto, variaveis, anos, args.uf, saida)
